#!/usr/bin/env python3
"""将 regime_era_calendar_matrix.csv 转为 .xlsx（UTF-8、冻结首行、首行加粗）。"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def main() -> None:
    if len(sys.argv) != 3:
        print("Usage: csv_matrix_to_xlsx.py <in.csv> <out.xlsx>", file=sys.stderr)
        sys.exit(2)
    in_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])

    wb = Workbook()
    ws = wb.active
    ws.title = "政权历法矩阵"

    with in_path.open(newline="", encoding="utf-8") as f:
        for row in csv.reader(f):
            ws.append(row)

    ws.freeze_panes = "A2"
    bold = Font(bold=True)
    for cell in ws[1]:
        cell.font = bold

    for i in range(1, ws.max_column + 1):
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = 10.0 if i == 1 else 36.0

    wb.save(out_path)


if __name__ == "__main__":
    main()
